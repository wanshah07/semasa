"""Wan's own reference library in OneDrive as a bank of post angles (Studio's nightly drafter, section 2, brought over
30 Sep 2026: "port all three before Friday").

`/40. HERMES/<folder>` holds one folder per domain (Cosmetic, Cosmetic Case Studies, Halal, Fatwa, Farmaseutikal, Food).
Studio's drafter read them and kept a topic bank of ANGLES and QUESTIONS, never client names, file numbers or content that
looks like a real application, and never a number as a verified fact. Here each file is read once, in the Composio
workbench (the download and the text extraction never leave it), the writer turns it into up to 8 angles, and each
angle is a row in the feed (`section = 'folder'`) that autofill and the page's "Jadikan idea" draw from. A row counts as
USED when an idea already carries its address (`onedrive://<file id>#<n>`), the same test the news feed uses.

Two kinds of file need their own handling:
  * a workbook with a sheet named "Post Matrix" (the LabMuffin/MoSkinLab matrix) is not mined by the writer: each ROW is
    already a post (column G citation, I English, J Malay, K visual, L design prompt), so the next few unused rows are
    put in the pool, ascending, and the pool is topped up as they are used;
  * a subfolder of `urgent post` named DDMMYY for today or tomorrow overrides that day: its files' rows carry
    `raw.urgent_for` and autofill gives them that day's slot ahead of the rota.

What the writer is sent is at most 12,000 characters of ONE file, with e-mail addresses, phone numbers and company names
masked first. It is the same kind of text Studio's routine read; the difference is that this writer is the gateway set in
Settings (rootsys / Mireld), not Claude. `sources.folders.use_writer: false` (SQL editor) stops the folders being mined.
"""

from __future__ import annotations

import json
import re
from datetime import UTC, datetime, timedelta
from typing import Any

from .log import get_logger

log = get_logger("semasa.folders")

DOMAINS = ("kosmetik", "makanan", "halal_my", "farmaseutikal", "fatwa", "sains_kosmetik", "kajian_kes", "lain")
MYT = timedelta(hours=8)
MAX_BYTES = 60_000_000
SEND_CHARS = 12_000
PER_CELL = 2          # files per workbench cell: 2 x 14,000 characters fits under the 36,000 the wrapper allows
ROOM = 18_000         # a cell stops opening files once its answer is this long
MATRIX_KIND = "Matriks siaran"

_LIST = '''
OUT = {"files": [], "errors": []}
def kids(path):
    res, err = call("ONE_DRIVE_LIST_FOLDER_CHILDREN", {"folder_path": path, "use_me_drive": True, "top": 100,
                    "select": ["id", "name", "size", "folder", "file", "lastModifiedDateTime", "webUrl"]}, P["account"])
    if err:
        OUT["errors"].append(path.rsplit("/", 1)[-1] + ": " + str(err)[:120])
        return []
    d = res.get("data") or {}
    if d.get("next_page_token"):
        OUT["errors"].append(path.rsplit("/", 1)[-1] + ": more than 100 items, only the first 100 read")
    return d.get("value") or []
def keep(v, folder, domain, urgent_for=None):
    if "folder" in v or not v.get("id"):         # the folder facet is only present on folders
        return
    OUT["files"].append({"id": v["id"], "name": v.get("name") or "", "size": v.get("size") or 0,
                         "mime": (v.get("file") or {}).get("mimeType") or "", "web": v.get("webUrl") or "",
                         "modified": (v.get("lastModifiedDateTime") or "")[:10], "folder": folder, "domain": domain,
                         "urgent_for": urgent_for})
for domain, folder in P["domains"].items():
    for v in kids(P["root"] + "/" + folder):
        keep(v, folder, domain)
if P.get("urgent"):
    for sub in kids(P["root"] + "/" + P["urgent"]):
        nm = str(sub.get("name") or "")
        if "folder" in sub and nm in P["urgent_dates"]:
            iso = "20" + nm[4:6] + "-" + nm[2:4] + "-" + nm[0:2]
            for v in kids(P["root"] + "/" + P["urgent"] + "/" + nm):
                keep(v, P["urgent"] + "/" + nm, None, iso)
'''

_READ = '''
import io, requests
OUT = {"docs": [], "skipped": []}
CAP = P["max_chars"]
def s(x, n):
    return ("" if x is None else str(x)).strip()[:n]
def pdf_text(b):
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(b)); n = len(r.pages)
    idx = list(range(min(n, 3))); rest = list(range(3, n))
    if len(rest) > 30:
        step = len(rest) / 27.0; rest = [rest[int(i * step)] for i in range(27)]
    txt = ""
    for i in idx + rest:
        t = (r.pages[i].extract_text() or "").strip()
        if t:
            txt += "\\n[hlm " + str(i + 1) + "]\\n" + t[:1200]
        if len(txt) >= CAP:
            break
    return {"kind": "pdf", "pages": n, "text": txt[:CAP]}
def docx_text(b):
    import docx
    d = docx.Document(io.BytesIO(b))
    return {"kind": "docx", "pages": 0, "text": "\\n".join(p.text for p in d.paragraphs if p.text.strip())[:CAP]}
def xlsx_read(b, m):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(b), read_only=True, data_only=True)
    names = wb.sheetnames
    pm = next((n for n in names if n.strip().lower() == "post matrix"), None)
    if pm:
        skip = set(m.get("skip") or []); rows = []
        for i, r in enumerate(wb[pm].iter_rows(values_only=True), 1):
            if i == 1 or i in skip:
                continue
            g = lambda k: (r[k] if len(r) > k else None)
            if not (s(g(8), 5) or s(g(9), 5)):
                continue
            rows.append({"row": i, "cite": s(g(6), 200), "en": s(g(8), 900), "bm": s(g(9), 900),
                         "visual": s(g(10), 250), "design": s(g(11), 250)})
            if len(rows) >= m.get("want", 8):
                break
        return {"kind": "matrix", "sheet": pm, "rows": rows}
    out = ""
    for n in names[:4]:
        out += "\\n[helaian " + n + "]\\n"
        for j, r in enumerate(wb[n].iter_rows(values_only=True)):
            if j >= 70:
                break
            line = " | ".join(s(c, 90) for c in (r or [])[:8] if c is not None and str(c).strip())
            if line:
                out += line + "\\n"
        if len(out) >= CAP:
            break
    return {"kind": "xlsx", "pages": 0, "text": out[:CAP]}
for f in P["files"]:
    if left() < 45 or len(json.dumps(OUT, ensure_ascii=False)) > P["room"]:
        OUT["skipped"].append(f["id"]); continue          # the answer must stay under what the workbench prints
    try:
        if f["size"] > P["max_bytes"]:
            OUT["docs"].append({"id": f["id"], "error": "too big (" + str(f["size"]) + " bytes)"}); continue
        dl, err = call("ONE_DRIVE_DOWNLOAD_FILE", {"item_id": f["id"], "file_name": f["name"]}, P["account"])
        url = ((dl.get("data") or {}).get("content") or {}).get("s3url")
        if err or not url:
            OUT["docs"].append({"id": f["id"], "error": "download: " + (str(err)[:100] or "no address")}); continue
        b = requests.get(url, timeout=110).content
        ext = f["name"].lower().rsplit(".", 1)[-1] if "." in f["name"] else ""
        if ext == "pdf":
            d = pdf_text(b)
        elif ext in ("docx",):
            d = docx_text(b)
        elif ext in ("xlsx", "xlsm"):
            d = xlsx_read(b, P.get("matrix") or {})
        else:
            OUT["docs"].append({"id": f["id"], "error": "type ." + ext + " is not read"}); continue
        OUT["docs"].append({"id": f["id"], **d})
    except Exception as e:
        OUT["docs"].append({"id": f["id"], "error": type(e).__name__ + ": " + str(e)[:120]})
'''

SYSTEM = """You mine a reference document from ws.regulab's own library for post ANGLES. ws.regulab is a Malaysian
regulatory consultancy (cosmetics, halal, food, pharmaceuticals); its readers are Malaysian SMEs. You are given the
folder's domain, the kind of document and text with [hlm N] page markers. Return up to {n} angles that a post could
build on. Each angle:
- "title": the question or wrong belief a Malaysian SME or consumer holds that this document answers, in Malaysian Malay
  (English regulatory terms are fine), at most 90 characters. Never Bahasa Indonesia.
- "summary": two sentences in Malaysian Malay, using ONLY the text. State a number, date or clause only when it is
  printed in the text, and say it is the document's own.
- "why": one sentence in Malaysian Malay: what a Malaysian business should do or know because of it.
- "cite": the document section, clause or page as printed (for example "hlm 12"), or "" when the text does not show one.
  Never invent a clause number.
- "domain": one of kosmetik, makanan, halal_my, farmaseutikal, fatwa, sains_kosmetik, kajian_kes, lain (default: the folder's).
Rules: take angles and questions only. NEVER a client's name, a company, a file or registration number, or anything that
looks like a real application or client case. Skip covers, contents pages, forewords and boilerplate. Fewer angles is
fine; none is fine.
Answer with ONE JSON object: {{"angles": [{{"title": "...", "summary": "...", "why": "...", "cite": "...", "domain": "..."}}]}}"""

_MASKS = (
    (re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+"), "[e-mel]"),
    (re.compile(r"(?<!\d)\+?\d[\d\s().-]{8,}\d(?!\d)"), "[nombor]"),
    (re.compile(r"\b(?:[A-Z][\w&.'-]+\s+){1,4}(?:Sdn\.?\s*Bhd\.?|Berhad|Bhd\.?)"), "[syarikat]"),
)


def mask(text: str) -> str:
    for rx, repl in _MASKS:
        text = rx.sub(repl, text)
    return text


def urgent_dates(now: datetime) -> list[str]:
    """Today's and tomorrow's date as DDMMYY, Malaysia time: the names of the subfolders that override a day."""
    d = (now + MYT).date()
    return [(d + timedelta(days=i)).strftime("%d%m%y") for i in (0, 1)]


def list_files(client: Any, cfg: dict[str, Any], now: datetime) -> dict[str, Any]:
    return client.cell(_LIST, {"root": cfg.get("root") or "/40. HERMES", "account": cfg.get("account") or None,
                               "domains": cfg.get("domains") or {}, "urgent": cfg.get("urgent") or "",
                               "urgent_dates": urgent_dates(now)}, thought="Semasa: what is in the OneDrive folders",
                       budget=100)


def read_files(client: Any, cfg: dict[str, Any], files: list[dict[str, Any]], matrix: dict[str, Any]) -> dict[str, Any]:
    """Read files in cells of PER_CELL. One cell may print about 36,000 characters, so a cell that read six files at
    14,000 characters each was refused whole and every file in it was lost (found live, 30 Sep 2026)."""
    slim = [{k: f[k] for k in ("id", "name", "size")} for f in files]
    docs: list[dict[str, Any]] = []
    skipped: list[str] = []
    for i in range(0, len(slim), PER_CELL):
        out = client.cell(_READ, {"account": cfg.get("account") or None, "files": slim[i:i + PER_CELL], "max_chars": 14_000,
                                  "max_bytes": MAX_BYTES, "matrix": matrix, "room": ROOM},
                          thought="Semasa: read new reference files (nothing is changed or uploaded)", budget=140) or {}
        docs += out.get("docs") or []
        skipped += out.get("skipped") or []
    return {"docs": docs, "skipped": skipped}


def kind_of(doc: dict[str, Any]) -> str:
    return {"pdf": "Dokumen PDF", "docx": "Dokumen Word", "xlsx": "Lembaran Excel"}.get(doc.get("kind"), "Dokumen")


def mine(llm: Any, doc: dict[str, Any], file: dict[str, Any], per_file: int) -> list[dict[str, Any]] | None:
    """Angles from one document; None when the writer gave no answer at all (so the file is tried again next run)."""
    if not llm or not getattr(llm, "configured", False):
        return None
    text = mask(str(doc.get("text") or ""))[:SEND_CHARS]
    if len(text.strip()) < 200:
        return []
    user = (f"Folder domain: {file.get('domain') or 'lain'}\nKind: {kind_of(doc)}"
            + (f", {doc['pages']} pages" if doc.get("pages") else "") + f"\n\nText:\n{text}")
    out = llm.chat_json(SYSTEM.format(n=per_file), user, max_tokens=420 * per_file + 200)
    if out is None:
        return None
    angles = []
    for a in (out or {}).get("angles") or []:
        title, summary = str(a.get("title") or "").strip(), str(a.get("summary") or "").strip()
        if title and summary:
            dom = str(a.get("domain") or "").strip().lower()
            angles.append({"title": title[:120], "summary": summary[:600], "why": str(a.get("why") or "").strip()[:300] or None,
                           "cite": str(a.get("cite") or "").strip()[:120],
                           "domain": dom if dom in DOMAINS else (file.get("domain") or "lain")})
    return angles[:per_file]


def angle_rows(file: dict[str, Any], doc: dict[str, Any], angles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """The feed rows for one mined file. An empty list of angles is recorded as one hidden marker, so a file with
    nothing in it is not read again every run."""
    base = {"section": "folder", "source": f"OneDrive · {file.get('folder') or ''}", "kind": kind_of(doc), "country": None,
            "lang": "ms", "summary_source": "llm", "published_at": None}
    raw0 = {"file_id": file["id"], "file": file.get("name"), "folder": file.get("folder"), "web_url": file.get("web") or None,
            **({"urgent_for": file["urgent_for"]} if file.get("urgent_for") else {})}
    if not angles:
        return [{**base, "title": "(tiada sudut dalam fail ini)", "url": f"onedrive://{file['id']}#0", "summary": None,
                 "why": None, "domain": file.get("domain"), "relevant": False, "raw": {**raw0, "empty": True}}]
    return [{**base, "title": a["title"], "url": f"onedrive://{file['id']}#{n}", "summary": a["summary"], "why": a["why"],
             "domain": a["domain"], "relevant": True, "raw": {**raw0, "cite": a["cite"]}}
            for n, a in enumerate(angles, 1)]


def matrix_rows(file: dict[str, Any], doc: dict[str, Any]) -> list[dict[str, Any]]:
    """Each matrix row is already a post: the English text is the summary, and everything else rides in `raw`."""
    out = []
    for r in doc.get("rows") or []:
        en = str(r.get("en") or r.get("bm") or "").strip()
        title = re.split(r"(?<=[.!?])\s", en, maxsplit=1)[0][:110]
        out.append({"section": "folder", "source": f"OneDrive · {file.get('folder') or ''}", "kind": MATRIX_KIND,
                    "country": None, "title": title or f"Baris {r['row']}", "url": f"onedrive://{file['id']}#r{r['row']}",
                    "summary": en[:600], "why": None, "domain": "kajian_kes", "relevant": True, "lang": "en",
                    "summary_source": "source", "published_at": None,
                    "raw": {"file_id": file["id"], "file": file.get("name"), "row": r["row"], "matrix": True,
                            "citation": r.get("cite"), "bm": r.get("bm"), "visual": r.get("visual"),
                            "design": r.get("design"), "web_url": file.get("web") or None}})
    return out


def file_id_of(url: str) -> str | None:
    m = re.match(r"^onedrive://([^#]+)#", url or "")
    return m.group(1) if m else None


def plan(files: list[dict[str, Any]], known: dict[str, Any], per_run: int) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """(files to mine now, workbooks that already feed the matrix pool). Newest first, so a file Wan just added is
    not queued behind an old one; `per_run` bounds the reading (a cell has 180 seconds)."""
    mined, matrix_ids = known["mined"], known["matrix_files"]
    fresh = sorted((f for f in files if f["id"] not in mined and f["id"] not in matrix_ids),
                   key=lambda f: (f.get("urgent_for") is None, f.get("modified") or ""), reverse=False)
    urgent = [f for f in fresh if f.get("urgent_for")]
    rest = sorted((f for f in fresh if not f.get("urgent_for")), key=lambda f: f.get("modified") or "", reverse=True)
    return (urgent + rest)[:per_run], [f for f in files if f["id"] in matrix_ids]


def now_utc() -> datetime:
    return datetime.now(UTC)


def dumps(x: Any) -> str:
    return json.dumps(x, ensure_ascii=False)
